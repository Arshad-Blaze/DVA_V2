"""Excel file discovery."""

from typing import List, Optional

import polars as pl

from dav_platform.core.contracts import ExcelSheetInfo


def discover_excel_sheets(file_path: str) -> List[ExcelSheetInfo]:
    """Discover Excel workbook structure.

    Returns sheet info for all sheets in the workbook.
    """
    try:
        import openpyxl
    except ImportError:
        return []

    try:
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        sheets = []

        for name in wb.sheetnames:
            ws = wb[name]
            rows = list(ws.iter_rows(max_row=10, values_only=True))
            if not rows:
                sheets.append(ExcelSheetInfo(name=name, confidence=0.0))
                continue

            # Detect header
            first_row = rows[0]
            has_header = _detect_header_from_values(first_row)

            # Get columns
            if has_header:
                columns = [str(v) if v else f"col_{i}" for i, v in enumerate(first_row)]
            else:
                columns = [f"col_{i}" for i in range(len(first_row))]

            # Estimate row count
            row_count = ws.max_row or 0
            col_count = ws.max_column or len(columns)

            # Confidence based on data presence
            data_rows = rows[1:] if has_header else rows
            non_empty = sum(1 for row in data_rows if any(v is not None for v in row))
            confidence = min(1.0, non_empty / max(len(data_rows), 1))

            sheets.append(ExcelSheetInfo(
                name=name,
                row_count=row_count,
                column_count=col_count,
                has_header=has_header,
                columns=columns,
                confidence=round(confidence, 2),
            ))

        wb.close()
        return sheets

    except Exception:
        return []


def select_candidate_sheet(sheets: List[ExcelSheetInfo]) -> Optional[str]:
    """Select the most likely data sheet.

    Heuristics:
    - Prefer sheets with more rows
    - Prefer sheets with headers
    - Prefer sheets with higher confidence
    """
    if not sheets:
        return None

    scored = []
    for sheet in sheets:
        score = 0.0
        score += min(1.0, sheet.row_count / 100) * 0.4  # Row count factor
        score += (0.3 if sheet.has_header else 0.0)  # Header bonus
        score += sheet.confidence * 0.3  # Confidence factor
        scored.append((sheet.name, score))

    scored.sort(key=lambda x: x[1], reverse=True)
    return scored[0][0] if scored else None


def _detect_header_from_values(row: tuple) -> bool:
    """Detect if a row contains headers based on value types."""
    if not row:
        return False

    alpha_count = 0
    for v in row:
        if v is None:
            continue
        s = str(v).strip()
        if s and any(c.isalpha() for c in s):
            alpha_count += 1

    return alpha_count >= len(row) / 2
